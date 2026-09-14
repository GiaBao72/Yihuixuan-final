import openpyxl
p = r'C:\Users\vgbvg\Downloads\yihuixuan_product_template.xlsx'
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb['Sản phẩm']
print('dims:', ws.dimensions, 'rows:', ws.max_row, 'cols:', ws.max_column)
print()
# header at row 3
hdr = [ws.cell(3,c).value for c in range(1, ws.max_column+1)]
for i,h in enumerate(hdr,1):
    print(f'  col{i}: {str(h)[:60]!r}')
print()
print('=== DATA ROWS 4..8 ===')
for r in range(4, 9):
    vals = [ws.cell(r,c).value for c in range(1, ws.max_column+1)]
    print(f'R{r}:')
    for i,v in enumerate(vals,1):
        s = str(v) if v is not None else ''
        if s:
            print(f'    c{i}: {s[:120]}')
    print()
wb.close()
